"""Plantilla de importación y ejemplo simulado reproducible (.xlsx).

    python -m app.plantilla        escribe datos/plantilla_salvamed.xlsx y datos/ejemplo_simulado.xlsx

El ejemplo usa la semilla 2026: con la misma fecha de generación produce los
mismos números. Sus 45 días de ventas terminan ayer, para que entren en la
ventana de 30 días completos. Todo el contenido del ejemplo es simulado.
"""

import io
import math
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.consumo import dia_local
from app.db import ahora

SEMILLA = 2026
DIAS_DE_HISTORIA = 45
INTRODUCIDO = PatternFill("solid", fgColor="E6F4F2")
SUPUESTO = PatternFill("solid", fgColor="FFF3C4")
TITULO = Font(bold=True, color="FFFFFF")
FONDO_TITULO = PatternFill("solid", fgColor="0F766E")

CATALOGO = (
    "nombre", "presentacion", "unidad", "sku", "clasificacion", "proveedor",
    "costo_compra_q", "precio_venta_q", "unidades_paquete", "existencia", "tiempo_entrega_dias", "vencimiento_proximo",
)
HISTORIAL = ("fecha", "medicamento", "presentacion", "cantidad")
CONFIGURACION = (
    ("dias_operacion_anio", 365, "Supuesto de demostración: la farmacia abre todos los días. Cámbialo si no es así."),
    ("nivel_servicio", 95, "Probabilidad de no agotarse durante el plazo de entrega (%)."),
    ("tasa_mantenimiento", 25, "Supuesto académico (% anual): H = costo de compra × tasa. No es un dato medido."),
    ("costo_pedido", 35, "S en quetzales. Escenario base: solo el envío."),
    ("tiempo_entrega", 3, "L en días. Escenario base: 3 días, constante."),
)

EJEMPLO = (
    # nombre, presentación, unidad, proveedor, costo, precio, paquete, existencia, L, consumo diario buscado
    ("Ceftriaxona", "1 g", "vial", "Proveedor A", 12.5, 16.9, 10, 60, 3, 6),
    ("Paracetamol", "1 g/100 ml", "frasco IV", "Proveedor A", 18.0, 24.3, 12, 150, 3, 15),
    ("Solución salina", "0.9 % · 100 ml", "bolsa", "Proveedor A", 5.5, 7.45, 20, 300, 3, 22),
    ("Dexametasona", "4 mg/ml · 2 ml", "ampolla", "Proveedor B", 5.9, 8.0, 10, 90, 3, 8),
    ("Metoclopramida", "10 mg/2 ml", "ampolla", "Proveedor B", 4.7, 6.35, 10, 100, 3, 7),
    ("Salbutamol", "2.5 mg/2.5 ml", "monodosis", "Proveedor B", 4.1, 5.55, 20, 120, 3, 10),
    ("Insulina cristalina", "U-100 · 10 ml", "vial", "Proveedor C", 88.0, 118.8, 1, 14, 4, 1.2),
    ("Adrenalina (epinefrina)", "1 mg/ml · 1 ml", "ampolla", "Proveedor C", 7.5, 10.15, 10, 20, 3, 1.5),
)


def _encabezado(hoja, columnas) -> None:
    hoja.append(list(columnas))
    for celda in hoja[1]:
        celda.font = TITULO
        celda.fill = FONDO_TITULO
    for i, nombre in enumerate(columnas, start=1):
        hoja.column_dimensions[hoja.cell(1, i).column_letter].width = max(14, len(nombre) + 4)


def _leeme(libro: Workbook, ejemplo: bool) -> None:
    hoja = libro.active
    hoja.title = "Léeme"
    hoja.column_dimensions["A"].width = 110
    lineas = [
        "SALVAMED · plantilla de importación" + (" · EJEMPLO SIMULADO" if ejemplo else ""),
        "",
        "Hojas:",
        "  Catálogo: un medicamento por fila. Solo nombre y unidad son obligatorios. La existencia de un medicamento nuevo entra como apertura.",
        "  Historial diario: ventas de días completos (una fila por día y medicamento). No mueve existencias: solo alimenta μD y σD.",
        "  Configuración: parámetros generales (A, p, tasa de mantenimiento, S, L).",
        "",
        "Colores: verde = dato que se introduce; amarillo = supuesto de demostración que conviene revisar.",
        "SALVAMED calcula lo demás (μD, σD, D, EOQ, R, IS, N, T, CT, H_IS). No escribas resultados en la plantilla.",
        "Un día observado sin ventas se escribe con cantidad 0; un día sin registro simplemente no se escribe.",
        "Al importar verás una vista previa con errores y procedencia antes de confirmar. Reimportar no duplica ni cambia existencias.",
    ]
    if ejemplo:
        lineas += [
            "",
            f"TODO ES SIMULADO: catálogo, costos, precios y existencias inventados; {DIAS_DE_HISTORIA} días de ventas GENERADOS con la semilla {SEMILLA}.",
            "Solo los nombres de los medicamentos están contextualizados en El Roble. No son datos entregados por El Roble.",
            "Los días terminan ayer respecto de la fecha en que se descargó este archivo.",
        ]
    for linea in lineas:
        hoja.append([linea])
    hoja["A1"].font = Font(bold=True, size=14)
    for fila in hoja.iter_rows():
        fila[0].alignment = Alignment(wrap_text=True)


def _configuracion(libro: Workbook) -> None:
    hoja = libro.create_sheet("Configuración")
    _encabezado(hoja, ("parametro", "valor", "nota"))
    for parametro, valor, nota in CONFIGURACION:
        hoja.append([parametro, valor, nota])
        hoja.cell(hoja.max_row, 2).fill = SUPUESTO
    hoja.column_dimensions["C"].width = 90


def _bytes(libro: Workbook) -> bytes:
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def plantilla() -> bytes:
    libro = Workbook()
    _leeme(libro, ejemplo=False)
    catalogo = libro.create_sheet("Catálogo")
    _encabezado(catalogo, CATALOGO)
    historial = libro.create_sheet("Historial diario")
    _encabezado(historial, HISTORIAL)
    _configuracion(libro)
    return _bytes(libro)


def _poisson(azar: random.Random, media: float) -> int:
    limite, k, p = math.exp(-media), 0, 1.0
    while True:
        p *= azar.random()
        if p <= limite:
            return k
        k += 1


def historia_simulada(hoy: date) -> list[tuple[date, str, str, int]]:
    """Ventas diarias generadas: Poisson con día de la semana y días altos o flojos al azar."""
    azar = random.Random(SEMILLA)
    patron = (1.1, 1.0, 1.0, 1.05, 1.2, 0.9, 0.65)
    filas = []
    for atras in range(DIAS_DE_HISTORIA, 0, -1):
        dia = hoy - timedelta(days=atras)
        factor = patron[dia.weekday()] * (0.5 if azar.random() < 0.08 else 1.5 if azar.random() > 0.92 else 1.0)
        for nombre, presentacion, *_, consumo in EJEMPLO:
            filas.append((dia, nombre, presentacion, _poisson(azar, consumo * factor)))
    return filas


def ejemplo(hoy: date | None = None) -> bytes:
    hoy = hoy or dia_local(ahora())
    libro = Workbook()
    _leeme(libro, ejemplo=True)
    catalogo = libro.create_sheet("Catálogo")
    _encabezado(catalogo, CATALOGO)
    for i, (nombre, presentacion, unidad, proveedor, costo, precio, paquete, existencia, plazo, _) in enumerate(EJEMPLO, start=1):
        catalogo.append(
            [nombre, presentacion, unidad, f"SIM-{i:03d}", "Regular", proveedor, costo, precio, paquete, existencia, plazo, None]
        )
        for columna in (7, 8, 10, 11):
            catalogo.cell(catalogo.max_row, columna).fill = SUPUESTO
    historial = libro.create_sheet("Historial generado")
    _encabezado(historial, HISTORIAL)
    for dia, nombre, presentacion, cantidad in historia_simulada(hoy):
        historial.append([dia.isoformat(), nombre, presentacion, cantidad])
    _configuracion(libro)
    return _bytes(libro)


if __name__ == "__main__":
    carpeta = Path(__file__).resolve().parent.parent / "datos"
    carpeta.mkdir(exist_ok=True)
    (carpeta / "plantilla_salvamed.xlsx").write_bytes(plantilla())
    (carpeta / "ejemplo_simulado.xlsx").write_bytes(ejemplo())
    print("Listo:", carpeta / "plantilla_salvamed.xlsx", carpeta / "ejemplo_simulado.xlsx")
