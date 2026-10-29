"""Importar catálogo, apertura, historia de consumo y configuración desde .xlsx, .csv o .tsv.

Flujo: subir → vista previa (mapeo de columnas, errores y procedencia) →
confirmar. Nada se guarda hasta confirmar.

Reglas para no duplicar ni cambiar existencias en silencio:
- Un medicamento se reconoce por nombre + presentación (sin tildes ni
  mayúsculas). Si ya existe, se actualizan sus datos de catálogo, pero su
  existencia NO cambia; la diferencia se informa. Solo si se marca la opción,
  se registra como ajuste con motivo.
- La existencia de un medicamento nuevo entra como movimiento de apertura.
- La historia de consumo no mueve existencias: solo alimenta μD y σD. Un día ya
  importado no se vuelve a importar, y no se importan días en que el sistema ya
  registra despachos (se contaría dos veces).
- Los totales semanales (sem_1 … sem_n) se guardan como semanas; no se
  convierten en consumos diarios observados. Se supone sem_1 la más antigua.
- Las columnas que el archivo ya trae calculadas (punto de reorden, stock de
  seguridad, pedido sugerido, …) no se importan: SALVAMED las recalcula.
"""

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app import inventario
from app.db import Configuracion, ConsumoHistorico, Importacion, Medicamento, Proveedor, ahora
from app.consumo import dia_local
from app.formulario import numero

CAMPOS_CATALOGO = {
    "nombre": ("Nombre", ("medicamento", "nombre", "producto")),
    "presentacion": ("Presentación", ("presentacion",)),
    "unidad": ("Unidad", ("unidad", "unidad_medida")),
    "sku": ("Código interno", ("sku", "codigo", "id_medicamento", "codigo_interno")),
    "clasificacion": ("Clasificación", ("clasificacion", "categoria")),
    "proveedor": ("Proveedor", ("proveedor", "proveedor_id")),
    "costo_compra": ("Costo de compra (Q)", ("costo_compra", "costo_compra_q", "costo_unitario_q")),
    "unidades_paquete": ("Unidades por paquete", ("unidades_paquete", "paquete")),
    "existencia": ("Existencia de apertura", ("existencia", "stock_actual", "existencia_inicial")),
    "tiempo_entrega": ("Plazo de entrega L (días)", ("tiempo_entrega", "tiempo_entrega_dias", "plazo")),
    "vencimiento_proximo": ("Vencimiento próximo", ("vencimiento_proximo", "vencimiento")),
}
CAMPOS_HISTORIAL = {
    "fecha": ("Fecha", ("fecha", "dia")),
    "medicamento": ("Medicamento", ("medicamento", "nombre")),
    "presentacion": ("Presentación", ("presentacion",)),
    "sku": ("Código interno", ("sku", "codigo", "id_medicamento")),
    "cantidad": ("Cantidad despachada", ("cantidad", "cantidad_despachada", "consumo", "despachos", "salidas", "unidades")),
}
CALCULADAS = {
    "consumo_semanal_promedio", "consumo_mensual_estimado", "stock_seguridad", "punto_reorden",
    "cobertura_dias", "revision_dias", "pedido_sugerido", "valor_inventario_q", "presupuesto_pedido_q",
    "estado", "dias_hasta_punto_reorden",
}
AMBIGUAS = {"precio_unitario_q": "No dice si es costo de compra: elige si corresponde."}
CONFIG = {
    "dias_operacion_anio": ("dias_operacion_anio", "a", "dias_operacion", "dias_habiles"),
    "nivel_servicio": ("nivel_servicio", "probabilidad_servicio", "p"),
    "tasa_mantenimiento": ("tasa_mantenimiento", "tasa_mantenimiento_anual"),
    "costo_pedido": ("costo_pedido", "s", "costo_envio"),
    "tiempo_entrega": ("tiempo_entrega", "l", "plazo_entrega"),
}


def clave(texto) -> str:
    texto = unicodedata.normalize("NFD", str(texto or "").strip().lower())
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", "_", texto).strip("_")


# ---------------------------------------------------------------- Lectura


@dataclass
class Hoja:
    nombre: str
    encabezados: list[str]
    filas: list[dict]
    tipo: str = ""
    mapeo: dict = field(default_factory=dict)


def _texto_celda(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _decodificar(contenido: bytes) -> tuple[str, str]:
    for codificacion in ("utf-8-sig", "cp1252"):
        try:
            return contenido.decode(codificacion), codificacion
        except UnicodeDecodeError:
            continue
    return contenido.decode("latin-1"), "latin-1"


def leer_archivo(nombre: str, contenido: bytes) -> tuple[list[Hoja], list[str]]:
    notas = []
    if nombre.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        hojas = []
        for ws in libro.worksheets:
            filas = [[_texto_celda(c) for c in fila] for fila in ws.iter_rows(values_only=True)]
            filas = [f for f in filas if any(f)]
            if len(filas) < 2:
                continue
            hojas.append(_hoja(ws.title, filas))
        return hojas, notas
    texto, codificacion = _decodificar(contenido)
    primera = texto.splitlines()[0] if texto else ""
    separador = max(("\t", ",", ";"), key=primera.count)
    nombre_sep = {"\t": "tabulaciones", ",": "comas", ";": "punto y coma"}[separador]
    notas.append(f"Texto en {codificacion}, columnas separadas por {nombre_sep}.")
    filas = [[c.strip() for c in fila] for fila in csv.reader(io.StringIO(texto), delimiter=separador)]
    filas = [f for f in filas if any(f)]
    if len(filas) < 2:
        return [], notas + ["El archivo no tiene filas de datos."]
    return [_hoja(nombre, filas)], notas


def _hoja(nombre: str, filas: list[list[str]]) -> Hoja:
    encabezados = [e or f"columna_{i + 1}" for i, e in enumerate(filas[0])]
    datos = [dict(zip(encabezados, fila + [""] * (len(encabezados) - len(fila)))) for fila in filas[1:]]
    hoja = Hoja(nombre, encabezados, datos)
    claves = {clave(e): e for e in encabezados}
    nombre_clave = clave(nombre)
    if "configuracion" in nombre_clave or {"parametro", "valor"} <= set(claves):
        hoja.tipo = "configuracion"
    elif "fecha" in claves and any(c in claves for c in ("cantidad", "cantidad_despachada", "consumo", "despachos", "salidas", "unidades")):
        hoja.tipo = "historial"
    elif any(c in claves for c in ("medicamento", "nombre", "producto")):
        hoja.tipo = "catalogo"
    else:
        hoja.tipo = "ignorar"
    campos = CAMPOS_CATALOGO if hoja.tipo == "catalogo" else CAMPOS_HISTORIAL if hoja.tipo == "historial" else {}
    for destino, (_, sinonimos) in campos.items():
        for sinonimo in sinonimos:
            if sinonimo in claves:
                hoja.mapeo[destino] = claves[sinonimo]
                break
    return hoja


def semanas_de(hoja: Hoja) -> list[str]:
    columnas = [(int(m.group(1)), e) for e in hoja.encabezados if (m := re.fullmatch(r"sem_?(\d+)", clave(e)))]
    return [e for _, e in sorted(columnas)]


# ---------------------------------------------------------------- Vista previa


def _num(texto: str, entero=False):
    if texto in ("", None):
        return None
    valor = numero(str(texto))
    if entero and not float(valor).is_integer():
        raise ValueError
    return int(valor) if entero else valor


def _fecha(texto: str) -> date | None:
    if not texto:
        return None
    texto = texto.strip()
    if re.fullmatch(r"\d{4}-\d{2}$", texto):
        return date.fromisoformat(texto + "-01")
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(texto[:10], formato).date()
        except ValueError:
            continue
    raise ValueError


def _existente(db: Session, nombre: str, presentacion: str, sku: str = "") -> Medicamento | None:
    candidatos = db.scalars(select(Medicamento).where(Medicamento.corrida_id.is_(None))).all()
    for m in candidatos:
        if clave(m.nombre) == clave(nombre) and clave(m.presentacion) == clave(presentacion):
            return m
    if sku:
        for m in candidatos:
            if m.sku and clave(m.sku) == clave(sku):
                return m
    return None


@dataclass
class Opciones:
    columna_ambigua: dict = field(default_factory=dict)  # columna → costo_compra | ignorar
    ajustar_existencias: bool = False
    aplicar_configuracion: bool = True
    importar_semanas: bool = True


def analizar(db: Session, hojas: list[Hoja], opciones: Opciones, archivo: str) -> dict:
    """Arma la vista previa: qué se crearía, qué se actualizaría y qué errores hay, fila por fila."""
    resultado = {"catalogo": [], "historial": [], "configuracion": [], "ignoradas": [], "avisos": []}
    for hoja in hojas:
        if hoja.tipo == "catalogo":
            _analizar_catalogo(db, hoja, opciones, archivo, resultado)
        elif hoja.tipo == "historial":
            _analizar_historial(db, hoja, archivo, resultado)
        elif hoja.tipo == "configuracion":
            _analizar_configuracion(hoja, archivo, resultado)
    return resultado


def _analizar_catalogo(db, hoja: Hoja, opciones: Opciones, archivo: str, resultado: dict) -> None:
    mapeo = dict(hoja.mapeo)
    for columna, destino in opciones.columna_ambigua.items():
        if columna in hoja.encabezados and destino == "costo_compra":
            mapeo[destino] = columna
    sin_uso = [e for e in hoja.encabezados if e not in mapeo.values() and e not in semanas_de(hoja)]
    for e in sin_uso:
        motivo = "calculada en el archivo: SALVAMED la recalcula" if clave(e) in CALCULADAS else AMBIGUAS.get(clave(e), "sin campo de destino")
        resultado["ignoradas"].append({"hoja": hoja.nombre, "columna": e, "motivo": motivo})
    semanas = semanas_de(hoja) if opciones.importar_semanas else []
    vistos = set()
    for numero_fila, fila in enumerate(hoja.filas, start=2):
        procedencia = f"{archivo} · {hoja.nombre} · fila {numero_fila}"
        datos, errores = {}, []
        for destino in CAMPOS_CATALOGO:
            columna = mapeo.get(destino)
            crudo = fila.get(columna, "") if columna else ""
            try:
                if destino in ("unidades_paquete", "existencia"):
                    datos[destino] = _num(crudo, entero=True)
                elif destino in ("costo_compra", "tiempo_entrega"):
                    datos[destino] = _num(crudo)
                elif destino == "vencimiento_proximo":
                    datos[destino] = _fecha(crudo)
                else:
                    datos[destino] = crudo.strip()
            except ValueError:
                errores.append(f"«{crudo}» no es válido para {CAMPOS_CATALOGO[destino][0]}")
        if not datos.get("nombre"):
            errores.append("Falta el nombre")
        if not datos.get("unidad"):
            errores.append("Falta la unidad")
        for campo in ("costo_compra", "tiempo_entrega"):
            if datos.get(campo) is not None and datos[campo] <= 0:
                errores.append(f"{CAMPOS_CATALOGO[campo][0]} debe ser mayor que 0")
        if datos.get("existencia") is not None and datos["existencia"] < 0:
            errores.append("La existencia no puede ser negativa")
        llave = (clave(datos.get("nombre")), clave(datos.get("presentacion")))
        if llave in vistos:
            errores.append("Fila repetida en el archivo (mismo nombre y presentación)")
        vistos.add(llave)
        semanales = []
        for columna in semanas:
            try:
                valor = _num(fila.get(columna, ""), entero=True)
            except ValueError:
                errores.append(f"«{fila.get(columna)}» no es un total semanal válido ({columna})")
                continue
            if valor is not None:
                semanales.append(valor)
        existente = None if errores else _existente(db, datos["nombre"], datos.get("presentacion", ""), datos.get("sku", ""))
        accion = "error" if errores else ("actualizar" if existente else "crear")
        aviso = ""
        if existente and datos.get("existencia") is not None and datos["existencia"] != existente.existencia:
            if opciones.ajustar_existencias:
                aviso = f"Se registrará un ajuste: {existente.existencia} → {datos['existencia']}"
            else:
                aviso = f"La existencia del sistema ({existente.existencia}) no cambia; el archivo dice {datos['existencia']}"
        if existente and semanales:
            ya = db.scalar(
                select(func.count()).select_from(ConsumoHistorico).where(
                    ConsumoHistorico.medicamento_id == existente.id, ConsumoHistorico.periodo == "semana"
                )
            )
            if ya:
                semanales = []
                aviso = (aviso + ". " if aviso else "") + "Ya tenía historia semanal: no se duplica"
        resultado["catalogo"].append(
            {
                "fila": numero_fila,
                "procedencia": procedencia,
                "accion": accion,
                "errores": errores,
                "aviso": aviso,
                "datos": {k: (v.isoformat() if isinstance(v, date) else v) for k, v in datos.items()},
                "semanas": semanales,
                "existente_id": existente.id if existente else None,
            }
        )


def _analizar_historial(db, hoja: Hoja, archivo: str, resultado: dict) -> None:
    mapeo = hoja.mapeo
    vistos = set()
    for numero_fila, fila in enumerate(hoja.filas, start=2):
        procedencia = f"{archivo} · {hoja.nombre} · fila {numero_fila}"
        errores = []
        nombre = fila.get(mapeo.get("medicamento", ""), "").strip()
        presentacion = fila.get(mapeo.get("presentacion", ""), "").strip()
        sku = fila.get(mapeo.get("sku", ""), "").strip()
        try:
            dia = _fecha(fila.get(mapeo.get("fecha", ""), ""))
            if dia is None:
                errores.append("Falta la fecha")
        except ValueError:
            dia = None
            errores.append(f"«{fila.get(mapeo.get('fecha', ''))}» no es una fecha válida")
        try:
            cantidad = _num(fila.get(mapeo.get("cantidad", ""), ""), entero=True)
            if cantidad is None or cantidad < 0:
                errores.append("La cantidad debe ser un entero de 0 o más")
        except ValueError:
            cantidad = None
            errores.append("La cantidad debe ser un entero")
        if dia and dia >= dia_local(ahora()):
            errores.append("Solo se importan días completos (anteriores a hoy)")
        resultado["historial"].append(
            {
                "fila": numero_fila,
                "procedencia": procedencia,
                "nombre": nombre,
                "presentacion": presentacion,
                "sku": sku,
                "fecha": dia.isoformat() if dia else "",
                "cantidad": cantidad,
                "errores": errores,
                "accion": "error" if errores else "importar",
            }
        )
        llave = (clave(nombre), clave(presentacion), dia)
        if not errores and llave in vistos:
            resultado["historial"][-1]["errores"].append("Día repetido en el archivo")
            resultado["historial"][-1]["accion"] = "error"
        vistos.add(llave)


def _analizar_configuracion(hoja: Hoja, archivo: str, resultado: dict) -> None:
    claves = {clave(e): e for e in hoja.encabezados}
    col_param = claves.get("parametro") or hoja.encabezados[0]
    col_valor = claves.get("valor") or hoja.encabezados[1]
    for numero_fila, fila in enumerate(hoja.filas, start=2):
        nombre = clave(fila.get(col_param, ""))
        destino = next((d for d, sin in CONFIG.items() if nombre in sin), None)
        if destino is None:
            continue
        try:
            valor = _num(fila.get(col_valor, ""))
        except ValueError:
            valor = None
        resultado["configuracion"].append(
            {
                "fila": numero_fila,
                "procedencia": f"{archivo} · {hoja.nombre} · fila {numero_fila}",
                "campo": destino,
                "valor": valor,
                "errores": [] if valor is not None else ["Valor no numérico"],
            }
        )


# ---------------------------------------------------------------- Aplicar


def _proveedor(db: Session, nombre: str, plazo: float | None) -> Proveedor | None:
    if not nombre:
        return None
    for p in db.scalars(select(Proveedor).where(Proveedor.corrida_id.is_(None))):
        if clave(p.nombre) == clave(nombre):
            return p
    p = Proveedor(nombre=nombre[:120], tiempo_entrega=None, costo_envio=None)
    db.add(p)
    db.flush()
    return p


def _semanas_inicio(hoy: date, total: int) -> list[date]:
    """La última semana termina ayer; sem_1 es la más antigua."""
    fin = hoy - timedelta(days=1)
    return [fin - timedelta(days=7 * (total - i) - 1) for i in range(total)]


def aplicar(db: Session, importacion: Importacion) -> dict:
    if importacion.estado != "vista_previa":
        raise inventario.ErrorDeOperacion("Esta importación ya se aplicó o se descartó.")
    vista = importacion.datos["vista"]
    momento = ahora()
    hoy = dia_local(momento)
    cuenta = {"creados": 0, "actualizados": 0, "ajustes": 0, "semanas": 0, "dias": 0, "dias_omitidos": 0, "configuracion": 0, "errores": 0}
    opciones = Opciones(**importacion.datos["opciones"])
    for fila in vista["catalogo"]:
        if fila["accion"] == "error":
            cuenta["errores"] += 1
            continue
        d = dict(fila["datos"])
        proveedor = _proveedor(db, d.pop("proveedor", ""), d.get("tiempo_entrega"))
        existencia = d.pop("existencia", None)
        if d.get("vencimiento_proximo"):
            d["vencimiento_proximo"] = date.fromisoformat(d["vencimiento_proximo"])
        if not d.get("unidades_paquete"):
            d["unidades_paquete"] = 1
        d = {k: v for k, v in d.items() if v not in (None, "") or k in ("presentacion",)}
        m = db.get(Medicamento, fila["existente_id"]) if fila["existente_id"] else None
        if m is None:
            m = inventario.dar_de_alta(db, d, existencia or 0, momento, usuario="Importación")
            m.origen = "importado"
            if proveedor:
                m.proveedor_id = proveedor.id
            cuenta["creados"] += 1
            for mv in db.new:
                if getattr(mv, "tipo", None) == "apertura" and getattr(mv, "medicamento_id", None) == m.id:
                    mv.nota = f"Apertura importada ({fila['procedencia']})"[:200]
        else:
            for campo, valor in d.items():
                setattr(m, campo, valor)
            if proveedor:
                m.proveedor_id = proveedor.id
            if m.costo_promedio is None:
                m.costo_promedio = m.costo_compra
            m.actualizado = momento
            cuenta["actualizados"] += 1
            if opciones.ajustar_existencias and existencia is not None and existencia != m.existencia:
                inventario._mover(
                    db, m, "ajuste", existencia - m.existencia, momento,
                    nota=f"Ajuste por importación ({fila['procedencia']})"[:200], usuario="Importación",
                )
                cuenta["ajustes"] += 1
        for inicio, cantidad in zip(_semanas_inicio(hoy, len(fila["semanas"])), fila["semanas"]):
            db.add(
                ConsumoHistorico(
                    medicamento_id=m.id, periodo="semana", inicio=inicio, cantidad=cantidad,
                    procedencia=fila["procedencia"][:200], importacion_id=importacion.id,
                )
            )
            cuenta["semanas"] += 1
        db.flush()
    for fila in vista["historial"]:
        if fila["accion"] == "error":
            cuenta["errores"] += 1
            continue
        m = _existente(db, fila["nombre"], fila["presentacion"], fila["sku"])
        if m is None:
            cuenta["dias_omitidos"] += 1
            continue
        dia = date.fromisoformat(fila["fecha"])
        repetido = db.scalar(
            select(func.count()).select_from(ConsumoHistorico).where(
                ConsumoHistorico.medicamento_id == m.id, ConsumoHistorico.periodo == "dia", ConsumoHistorico.inicio == dia
            )
        )
        if repetido or (m.observado_desde and dia >= m.observado_desde):
            cuenta["dias_omitidos"] += 1
            continue
        db.add(
            ConsumoHistorico(
                medicamento_id=m.id, periodo="dia", inicio=dia, cantidad=fila["cantidad"],
                procedencia=fila["procedencia"][:200], importacion_id=importacion.id,
                generado="generad" in fila["procedencia"].lower(),
            )
        )
        db.flush()
        cuenta["dias"] += 1
    if opciones.aplicar_configuracion:
        cfg = db.get(Configuracion, 1)
        for fila in vista["configuracion"]:
            if fila["errores"]:
                continue
            valor = fila["valor"]
            if fila["campo"] == "tasa_mantenimiento" and valor > 1:
                valor = valor / 100
            if fila["campo"] == "dias_operacion_anio":
                valor = int(valor)
            setattr(cfg, fila["campo"], valor)
            cuenta["configuracion"] += 1
    importacion.estado = "aplicada"
    importacion.aplicado_en = momento
    importacion.resumen = cuenta
    db.commit()
    return cuenta


def preparar(db: Session, archivo: str, contenido: bytes, opciones: Opciones) -> Importacion:
    hojas, notas = leer_archivo(archivo, contenido)
    vista = analizar(db, hojas, opciones, archivo)
    vista["avisos"] = notas + vista["avisos"]
    importacion = Importacion(
        archivo=archivo[:200],
        datos={
            "hojas": [
                {"nombre": h.nombre, "tipo": h.tipo, "encabezados": h.encabezados, "mapeo": h.mapeo, "filas": h.filas, "semanas": semanas_de(h)}
                for h in hojas
            ],
            "opciones": opciones.__dict__,
            "vista": vista,
        },
        creado=ahora(),
    )
    db.add(importacion)
    db.commit()
    return importacion


def reanalizar(db: Session, importacion: Importacion, opciones: Opciones, mapeos: dict[str, dict]) -> None:
    """Vuelve a armar la vista previa con otro mapeo de columnas u otras opciones."""
    hojas = []
    for h in importacion.datos["hojas"]:
        hoja = Hoja(h["nombre"], h["encabezados"], h["filas"], h["tipo"], dict(h["mapeo"]))
        for destino, columna in mapeos.get(h["nombre"], {}).items():
            if columna in hoja.encabezados:
                hoja.mapeo[destino] = columna
            elif not columna:
                hoja.mapeo.pop(destino, None)
        hojas.append(hoja)
    vista = analizar(db, hojas, opciones, importacion.archivo)
    vista["avisos"] = importacion.datos["vista"].get("avisos", [])
    datos = dict(importacion.datos)
    datos["hojas"] = [dict(h, mapeo=hoja.mapeo) for h, hoja in zip(importacion.datos["hojas"], hojas)]
    datos["opciones"] = opciones.__dict__
    datos["vista"] = vista
    importacion.datos = datos
    db.commit()
